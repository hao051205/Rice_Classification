"""
build_perfect_presentation.py - Tạo bộ Slide báo cáo 12 trang TỐI ƯU TRỰC QUAN (Visual, Ít Chữ)
Kế thừa mẫu nền chính thức UTEHY từ file BAI14_Bảo mật trong Truyền thông IoT (1).pptx.
Giữ lại thông tin 2 thành viên (Chử Quốc Hào & Vũ Đình Hiếu).
Thiết kế lại toàn diện: Ít chữ, nhiều hình ảnh, thẻ chỉ số KPI to rõ, sơ đồ quy trình trực quan và Speaker Notes chi tiết.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE_PATH = os.path.join(BASE_DIR, "slides", "BAI14_Bảo mật trong Truyền thông IoT (1).pptx")
SLIDES_PATH = os.path.join(BASE_DIR, "slides", "presentation.pptx")
ALT_PATH = os.path.join(BASE_DIR, "slides", "presentation_final.pptx")
FIGURES_DIR = os.path.join(BASE_DIR, "report", "figures")

# Khởi tạo presentation từ template UTEHY gốc
prs = Presentation(TEMPLATE_PATH)

# Xóa toàn bộ 51 slides ban đầu của bài giảng IoT, giữ nguyên masters và layouts
while len(prs.slides._sldIdLst) > 0:
    sldId = prs.slides._sldIdLst[0]
    rId = sldId.rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[0]

# Bảng màu chuẩn UTEHY và phong cách đồ họa học thuật hiện đại
C_UTEHY_GREEN = RGBColor(0, 168, 89)    # #00A859 (Xanh lá UTEHY)
C_UTEHY_GOLD = RGBColor(255, 192, 0)    # #FFC000 (Vàng UTEHY)
C_NAVY_DARK = RGBColor(16, 42, 67)      # #102A43 (Tiêu đề đậm)
C_DARK = RGBColor(33, 37, 41)           # #212529 (Văn bản nội dung)
C_MUTED = RGBColor(100, 116, 139)       # #64748B (Chú thích xám)
C_CARD_BG = RGBColor(255, 255, 255)     # Nền card trắng tinh tế
C_CARD_BORDER = RGBColor(226, 232, 240) # Viền card mảnh
C_WHITE = RGBColor(255, 255, 255)       # Trắng
C_GREEN_BG = RGBColor(240, 253, 244)    # Xanh lá nhạt (Highlight tích cực)
C_RED = RGBColor(185, 28, 28)           # Đỏ cảnh báo
C_RED_BG = RGBColor(254, 242, 242)      # Đỏ nhạt (Card cảnh báo)
C_BLUE_LIGHT = RGBColor(239, 246, 255)  # Xanh dương nhạt
C_GOLD_BG = RGBColor(254, 243, 199)     # Vàng nhạt (Card lưu ý)


def create_content_slide():
    """Tạo slide nội dung dùng layout 1 (kế thừa nền UTEHY image8.png)."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    for shp in list(slide.shapes):
        sp = shp._element
        sp.getparent().remove(sp)
    return slide


def add_slide_header(slide, title_text):
    """
    Đặt tiêu đề slide màu trắng vào bên trong dải băng màu xanh lá cây của UTEHY ở đầu slide.
    Phần logo và tên trường bên phải (x > 10.2 inch) được bảo toàn tuyệt đối.
    """
    tb_title = slide.shapes.add_textbox(Inches(0.6), Inches(0.12), Inches(9.4), Inches(0.64))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0

    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(20)
    p_title.font.bold = True
    p_title.font.color.rgb = C_WHITE


def add_slide_footer(slide, current_page, total_pages=12):
    """Footer đặt ở vùng an toàn không che phần họa tiết vàng góc phải."""
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.90), Inches(10.5), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = C_CARD_BORDER
    line.line.color.rgb = C_CARD_BORDER

    tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(7.5), Inches(0.35))
    tf_l = tb_l.text_frame
    p_l = tf_l.paragraphs[0]
    p_l.text = "Học máy cơ bản — Báo cáo bảo vệ Project 17 | GVHD: PGS.TS. Nguyễn Văn Hậu"
    p_l.font.size = Pt(10)
    p_l.font.color.rgb = C_MUTED

    tb_r = slide.shapes.add_textbox(Inches(9.0), Inches(6.95), Inches(2.2), Inches(0.35))
    tf_r = tb_r.text_frame
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.text = f"Slide {current_page} / {total_pages}"
    p_r.font.size = Pt(10.5)
    p_r.font.bold = True
    p_r.font.color.rgb = C_UTEHY_GREEN


def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, shape=MSO_SHAPE.RECTANGLE):
    """Thêm một khối khung (card) đồ họa chuyên nghiệp."""
    card = slide.shapes.add_shape(shape, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card


def add_kpi_badge(slide, left, top, width, height, num_str, label_str, bg_color=C_WHITE, num_color=C_UTEHY_GREEN, border_color=C_CARD_BORDER):
    """Thêm khối hiển thị chỉ số nổi bật (KPI Block) trực quan."""
    add_card(slide, left, top, width, height, bg_color=bg_color, border_color=border_color, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tb = slide.shapes.add_textbox(left, top + Inches(0.08), width, height - Inches(0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    p_num = tf.paragraphs[0]
    p_num.text = num_str
    p_num.font.size = Pt(22)
    p_num.font.bold = True
    p_num.font.color.rgb = num_color
    p_num.alignment = PP_ALIGN.CENTER
    
    p_lbl = tf.add_paragraph()
    p_lbl.text = label_str
    p_lbl.font.size = Pt(10)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = C_NAVY_DARK
    p_lbl.alignment = PP_ALIGN.CENTER
    p_lbl.space_before = Pt(2)


def set_speaker_notes(slide, notes_text):
    """Gắn lời thoại thuyết trình chuẩn mực vào phần Speaker Notes."""
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text.strip()


# ============================================================
# SLIDE 1: TIÊU ĐỀ & THÔNG TIN DỰ ÁN (Sử dụng Title Layout 0)
# ============================================================
s1 = prs.slides.add_slide(prs.slide_layouts[0])
for shp in list(s1.shapes):
    sp = shp._element
    sp.getparent().remove(sp)

# Badge Môn học (Theo đúng chỉnh sửa của người dùng)
tb_sub = s1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.5))
tf_sub = tb_sub.text_frame
p = tf_sub.paragraphs[0]
p.text = "MÔN HỌC: HỌC MÁY CƠ BẢN"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN

# Tiêu đề Project Card
add_card(s1, Inches(0.8), Inches(2.1), Inches(11.733), Inches(2.5), bg_color=C_WHITE, border_color=C_UTEHY_GREEN)
tb_title = s1.shapes.add_textbox(Inches(1.1), Inches(2.25), Inches(11.133), Inches(2.2))
tf_title = tb_title.text_frame
tf_title.word_wrap = True

p = tf_title.paragraphs[0]
p.text = "PROJECT 17: PHÂN LOẠI HAI GIỐNG LÚA GẠO CAMMEO VÀ OSMANCIK\nBẰNG CÂY QUYẾT ĐỊNH VÀ RỪNG NGẪU NHIÊN"
p.font.size = Pt(24)
p.font.bold = True
p.font.color.rgb = C_NAVY_DARK
p.alignment = PP_ALIGN.CENTER

p = tf_title.add_paragraph()
p.text = "Nghiên cứu ứng dụng Cost-Complexity Pruning (ccp_alpha), Random Forest và triển khai Web Service"
p.font.size = Pt(13)
p.font.italic = True
p.font.color.rgb = C_MUTED
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(8)

# Thông tin GVHD (Trái)
add_card(s1, Inches(0.8), Inches(4.8), Inches(5.7), Inches(1.9), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_sup = s1.shapes.add_textbox(Inches(1.0), Inches(4.95), Inches(5.3), Inches(1.6))
tf_sup = tb_sup.text_frame
p = tf_sup.paragraphs[0]
p.text = "GIẢNG VIÊN HƯỚNG DẪN:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN

p = tf_sup.add_paragraph()
p.text = "PGS. TS. Nguyễn Văn Hậu"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = C_NAVY_DARK
p.space_before = Pt(3)

p = tf_sup.add_paragraph()
p.text = "Khoa Công nghệ thông tin — Trường ĐH Sư phạm Kỹ thuật Hưng Yên"
p.font.size = Pt(10.5)
p.font.color.rgb = C_MUTED
p.space_before = Pt(2)

# Thông tin Sinh viên (Phải - Giữ trọn vẹn cả 2 sinh viên đã bổ sung)
add_card(s1, Inches(6.833), Inches(4.8), Inches(5.7), Inches(1.9), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_stu = s1.shapes.add_textbox(Inches(7.033), Inches(4.95), Inches(5.3), Inches(1.6))
tf_stu = tb_stu.text_frame
p = tf_stu.paragraphs[0]
p.text = "SINH VIÊN THỰC HIỆN:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN

p = tf_stu.add_paragraph()
p.text = "1. Chử Quốc Hào  —  MSV: 10123112"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = C_NAVY_DARK
p.space_before = Pt(4)

p = tf_stu.add_paragraph()
p.text = "2. Vũ Đình Hiếu     —  MSV: 10123131"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = C_NAVY_DARK
p.space_before = Pt(2)

p = tf_stu.add_paragraph()
p.text = "Lớp: 10123O.1 — Khoa Công nghệ thông tin"
p.font.size = Pt(10.5)
p.font.color.rgb = C_MUTED
p.space_before = Pt(3)

set_speaker_notes(s1, """
Kính thưa Thầy giáo PGS.TS. Nguyễn Văn Hậu và toàn thể Hội đồng chấm bảo vệ! Hôm nay, nhóm sinh viên chúng em gồm Chử Quốc Hào và Vũ Đình Hiếu xin phép được báo cáo đề tài Bài tập lớn môn Học máy cơ bản: 'Phân loại hai giống lúa gạo Cammeo và Osmancik bằng Cây quyết định và Rừng ngẫu nhiên'. Đề tài tập trung giải quyết bài toán tự động hóa kiểm định nông sản thông qua quy trình thực nghiệm nghiêm ngặt, tối ưu cây bằng Cost-Complexity Pruning và đóng gói thành Web Service hoàn chỉnh.
""")


# ============================================================
# SLIDE 2: ĐẶT VẤN ĐỀ / ĐỘNG LỰC THỰC TIỄN (Visual Cards + Hình 1)
# ============================================================
s2 = create_content_slide()
add_slide_header(s2, "Giám Định Thủ Công Làm Chậm Dây Chuyền Và Bộc Lộ Rủi Ro Lẫn Tạp")
add_slide_footer(s2, 2)

# Cột trái: 3 Khối đồ họa tóm lược ngắn gọn, trực quan (Không viết văn dài)
cards_s2 = [
    (Inches(0.8), Inches(1.15), Inches(5.7), Inches(1.7), "THƯƠNG PHẨM & GIÁ TRỊ",
     [("Cammeo (Hạt dài, giá cao)", "Osmancik (Hạt tròn, phổ thông)."),
      ("Lẫn tạp giống:", "Làm giảm phẩm cấp xuất khẩu, tổn hại uy tín thương hiệu.")]),
    (Inches(0.8), Inches(3.05), Inches(5.7), Inches(1.7), "HẠN CHẾ GIÁM ĐỊNH THỦ CÔNG",
     [("Phân loại bằng mắt:", "Chậm, mỏi mắt sau ca dài, mang tính chủ quan."),
      ("Bất khả thi quy mô lớn:", "Không đáp ứng được công suất hàng tấn/giờ trên băng chuyền.")]),
    (Inches(0.8), Inches(4.95), Inches(5.7), Inches(1.75), "GIẢI PHÁP THỊ GIÁC MÁY & AI",
     [("Hệ thống chụp ảnh:", "Đo đạc tự động 7 chỉ số hình học vi mô của từng hạt gạo."),
      ("Mô hình Máy học:", "Phân loại tức thì (< 1 ms), điều khiển van khí nén đạt độ chuẩn > 90%.")])
]

for left, top, w, h, header, items in cards_s2:
    add_card(s2, left, top, w, h, bg_color=C_WHITE, border_color=C_CARD_BORDER)
    tb = s2.shapes.add_textbox(left + Inches(0.2), top + Inches(0.12), w - Inches(0.4), h - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = header
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = C_UTEHY_GREEN
    for b_title, b_desc in items:
        p_item = tf.add_paragraph()
        p_item.text = f"• {b_title} "
        p_item.font.bold = True
        p_item.font.size = Pt(10.5)
        p_item.font.color.rgb = C_NAVY_DARK
        p_item.space_before = Pt(3)
        run = p_item.add_run()
        run.text = b_desc
        run.font.bold = False
        run.font.size = Pt(10)
        run.font.color.rgb = C_DARK

# Cột phải: Hình 1 Phân bố đặc trưng
img_p = os.path.join(FIGURES_DIR, "01_distributions.png")
if os.path.exists(img_p):
    s2.shapes.add_picture(img_p, Inches(6.75), Inches(1.15), width=Inches(5.78))
    tb_cap = s2.shapes.add_textbox(Inches(6.75), Inches(6.25), Inches(5.78), Inches(0.35))
    tb_cap.text_frame.paragraphs[0].text = "Hình 1: Phân bố 7 đặc trưng hình học giữa Cammeo (xanh) và Osmancik (cam)"
    tb_cap.text_frame.paragraphs[0].font.size = Pt(9.5)
    tb_cap.text_frame.paragraphs[0].font.italic = True
    tb_cap.text_frame.paragraphs[0].font.color.rgb = C_MUTED
    tb_cap.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

set_speaker_notes(s2, """
Trong xuất khẩu lúa gạo, Cammeo và Osmancik có giá thành và thị trường khác biệt. Nếu phân loại bằng mắt, công nhân nhanh chóng mệt mỏi dẫn đến nhầm lẫn, gây lẫn tạp lô hàng. Hệ thống chụp ảnh trên băng chuyền kết hợp máy học giúp đo đạc các kích thước vi mô và ra quyết định chính xác chỉ trong vài mili-giây.
""")


# ============================================================
# SLIDE 3: PHÁT BIỂU BÀI TOÁN & BIỂU ĐỒ PHÂN TÁN 2D (Visualized)
# ============================================================
s3 = create_content_slide()
add_slide_header(s3, "Phát Biểu Bài Toán Phân Loại Nhị Phân Dựa Trên 7 Đặc Trưng Hình Học")
add_slide_footer(s3, 3)

# Cột trái: 3 Khối Card thông số trực quan + Toán học
add_card(s3, Inches(0.8), Inches(1.15), Inches(5.7), Inches(1.65), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_m1 = s3.shapes.add_textbox(Inches(1.0), Inches(1.25), Inches(5.3), Inches(1.45))
tf_m1 = tb_m1.text_frame
tf_m1.word_wrap = True
p = tf_m1.paragraphs[0]
p.text = "MÔ HÌNH HÓA TOÁN HỌC"
p.font.size = Pt(11.5)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN
p1 = tf_m1.add_paragraph()
p1.text = "• Đầu vào: Véc-tơ số thực 7 chiều:  x in R^7"
p1.font.bold = True
p1.font.size = Pt(11)
p1.font.color.rgb = C_NAVY_DARK
p1.space_before = Pt(3)
p2 = tf_m1.add_paragraph()
p2.text = "• Đầu ra: Nhãn nhị phân:  y in {Cammeo, Osmancik}"
p2.font.bold = True
p2.font.size = Pt(11)
p2.font.color.rgb = C_NAVY_DARK
p2.space_before = Pt(2)
p3 = tf_m1.add_paragraph()
p3.text = "• Hàm mục tiêu: Ước lượng xác suất P(y = Osmancik | x)"
p3.font.size = Pt(10)
p3.font.color.rgb = C_MUTED
p3.space_before = Pt(2)

add_card(s3, Inches(0.8), Inches(2.95), Inches(5.7), Inches(1.75), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_m2 = s3.shapes.add_textbox(Inches(1.0), Inches(3.05), Inches(5.3), Inches(1.55))
tf_m2 = tb_m2.text_frame
tf_m2.word_wrap = True
p = tf_m2.paragraphs[0]
p.text = "QUY MÔ DỮ LIỆU & 7 ĐẶC TRƯNG HÌNH THÁI"
p.font.size = Pt(11.5)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN
p1 = tf_m2.add_paragraph()
p1.text = "• 3.810 hạt chuẩn từ UCI: Osmancik 57,2% (2.180) | Cammeo 42,8% (1.630)"
p1.font.size = Pt(10)
p1.font.color.rgb = C_DARK
p1.space_before = Pt(3)
p2 = tf_m2.add_paragraph()
p2.text = "• Đầy đủ 7 đặc trưng:"
p2.font.bold = True
p2.font.size = Pt(10)
p2.font.color.rgb = C_NAVY_DARK
p2.space_before = Pt(3)
p3 = tf_m2.add_paragraph()
p3.text = "  Area, Perimeter, Major_Axis_Length, Minor_Axis_Length, Eccentricity, Convex_Area, Extent."
p3.font.size = Pt(10)
p3.font.color.rgb = C_DARK

# 2 KPI Mini Cards ở dưới bên trái
add_kpi_badge(s3, Inches(0.8), Inches(4.85), Inches(2.75), Inches(1.85),
              "57,22%", "BASELINE TỐI THIỂU\n(Dummy: Most Frequent)",
              bg_color=C_RED_BG, num_color=C_RED, border_color=C_RED)
add_kpi_badge(s3, Inches(3.75), Inches(4.85), Inches(2.75), Inches(1.85),
              "> 90,0%", "MỤC TIÊU ACCURACY & F1\n(Test Set độc lập)",
              bg_color=C_GREEN_BG, num_color=C_UTEHY_GREEN, border_color=C_UTEHY_GREEN)

# Cột phải: Nhúng biểu đồ phân tán 2D (Major vs Minor Axis) để trực quan hóa không gian đặc trưng
img_scatter = os.path.join(FIGURES_DIR, "04_scatter_major_minor.png")
if os.path.exists(img_scatter):
    s3.shapes.add_picture(img_scatter, Inches(6.75), Inches(1.15), width=Inches(5.78))
    tb_cap = s3.shapes.add_textbox(Inches(6.75), Inches(6.25), Inches(5.78), Inches(0.35))
    tb_cap.text_frame.paragraphs[0].text = "Hình 1b: Không gian phân tách 2D giữa Major_Axis và Minor_Axis của 3.810 hạt gạo"
    tb_cap.text_frame.paragraphs[0].font.size = Pt(9.5)
    tb_cap.text_frame.paragraphs[0].font.italic = True
    tb_cap.text_frame.paragraphs[0].font.color.rgb = C_MUTED
    tb_cap.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

set_speaker_notes(s3, """
Về mặt toán học, bài toán được phát biểu là: Cho véc-tơ đầu vào x thuộc R^7 gồm đầy đủ 7 đặc trưng hình thái: Area, Perimeter, Major_Axis_Length, Minor_Axis_Length, Eccentricity, Convex_Area và Extent. Mô hình cần học hàm ánh xạ để dự đoán nhãn nhị phân y thuộc {Cammeo, Osmancik}. Dữ liệu gồm 3.810 mẫu hạt gạo chuẩn từ UCI. Quan sát biểu đồ phân tán 2D ở Hình 1b, hai giống lúa có sự phân tách tương đối nhưng vẫn tồn tại dải ranh giới giao thoa tự nhiên. Mục tiêu là mô hình phải vượt mốc ngẫu nhiên 57,22% và đạt độ chính xác trên 90% trên tập Test.
""")


# ============================================================
# SLIDE 4: THÁCH THỨC DỮ LIỆU & RỦI RO RÒ RỈ (Alert Cards + Hình 2)
# ============================================================
s4 = create_content_slide()
add_slide_header(s4, "Đa Cộng Tuyến Mạnh Và Rủi Ro Rò Rỉ Dữ Liệu Trong Huấn Luyện")
add_slide_footer(s4, 4)

# Cột trái: 3 Khối Card phân tích ngắn gọn, dùng Badge làm nổi bật
cards_s4 = [
    (Inches(0.8), Inches(1.15), Inches(5.7), Inches(1.75), "ĐA CỘNG TUYẾN CAO Ở 4 BIẾN KÍCH CỠ",
     [("Nhóm 4 biến:", "Area, Perimeter, Major_Axis, Convex_Area có r >= 0.90."),
      ("Nhiều cặp cực cao:", "Area - Convex_Area (r = 1.00), Area - Perimeter (r = 0.97)."),
      ("Hạn chế tuyến tính:", "Ma trận hiệp phương sai suy biến, trọng số hồi quy bị méo mó.")]),
    (Inches(0.8), Inches(3.05), Inches(5.7), Inches(1.75), "RỦI RO RÒ RỈ DỮ LIỆU (DATA LEAKAGE)",
     [("Rủi ro tính toán sớm:", "Nếu tính min, max, mean hoặc chuẩn hóa trước khi chia tập..."),
      ("Hậu quả nghiêm trọng:", "Thông tin tập Test sẽ ngấm vào Train, gây hiện tượng học vẹt."),
      ("Độ chính xác ảo:", "Điểm số trên thực nghiệm rất cao nhưng thất bại hoàn toàn trên thực tế.")]),
    (Inches(0.8), Inches(4.95), Inches(5.7), Inches(1.75), "THÁCH THỨC CÂY TỰ DO & CÂU HỎI NGHIÊN CỨU",
     [("Cây tự do (CART):", "Miễn nhiễm đa cộng tuyến nhưng rất dễ chia nhánh sâu, Overfitting nặng."),
      ("Câu hỏi nghiên cứu:", "Cắt tỉa ccp_alpha có tạo ra cây tối giản đủ tốt để giải bài toán minh bạch?")])
]

for left, top, w, h, header, items in cards_s4:
    add_card(s4, left, top, w, h, bg_color=C_WHITE, border_color=C_CARD_BORDER)
    tb = s4.shapes.add_textbox(left + Inches(0.2), top + Inches(0.12), w - Inches(0.4), h - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = header
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = C_RED if "RÒ RỈ" in header else C_UTEHY_GREEN
    for b_title, b_desc in items:
        p_item = tf.add_paragraph()
        p_item.text = f"• {b_title} "
        p_item.font.bold = True
        p_item.font.size = Pt(10.5)
        p_item.font.color.rgb = C_NAVY_DARK
        p_item.space_before = Pt(3)
        run = p_item.add_run()
        run.text = b_desc
        run.font.bold = False
        run.font.size = Pt(10)
        run.font.color.rgb = C_DARK

# Cột phải: Hình 2 Ma trận tương quan
img_p = os.path.join(FIGURES_DIR, "03_correlation_matrix.png")
if os.path.exists(img_p):
    s4.shapes.add_picture(img_p, Inches(6.75), Inches(1.15), width=Inches(5.78))
    tb_cap = s4.shapes.add_textbox(Inches(6.75), Inches(6.25), Inches(5.78), Inches(0.35))
    tb_cap.text_frame.paragraphs[0].text = "Hình 2: Ma trận tương quan Pearson khẳng định hiện tượng đa cộng tuyến (r >= 0.90, nhiều cặp r >= 0.97)"
    tb_cap.text_frame.paragraphs[0].font.size = Pt(9.5)
    tb_cap.text_frame.paragraphs[0].font.italic = True
    tb_cap.text_frame.paragraphs[0].font.color.rgb = C_MUTED
    tb_cap.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

set_speaker_notes(s4, """
Phân tích ma trận tương quan ở Hình 2 cho thấy rõ hiện tượng đa cộng tuyến: nhóm 4 biến kích cỡ đều có r >= 0.90, trong đó nhiều cặp r >= 0.97. Ngoài ra, một rủi ro cực lớn trong thực hành học máy là Rò rỉ dữ liệu (Data Leakage): nếu chuẩn hóa dữ liệu trên toàn bộ tập trước khi chia, mô hình sẽ 'biết trước' thông tin tập Test. Cây quyết định tự do giải quyết được đa cộng tuyến nhưng lại dễ bị quá khớp. Do đó, nghiên cứu đặt trọng tâm vào cắt tỉa ccp_alpha để tìm ra cây tối giản.
""")


# ============================================================
# SLIDE 5: THIẾT LẬP THỰC NGHIỆM & QUY TRÌNH CHỐNG RÒ RỈ (Flowchart Trực quan)
# ============================================================
s5 = create_content_slide()
add_slide_header(s5, "Thiết Lập Quy Trình Stratified 60/20/20 Và Cột Mốc Dummy Baseline")
add_slide_footer(s5, 5)

# Cột 1: Sơ đồ phân chia Stratified 60/20/20
add_card(s5, Inches(0.8), Inches(1.15), Inches(3.64), Inches(5.55), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_s5_1 = s5.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(3.24), Inches(5.25))
tf_s5_1 = tb_s5_1.text_frame
tf_s5_1.word_wrap = True
p = tf_s5_1.paragraphs[0]
p.text = "1. PHÂN CHIA STRATIFIED"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN

flow_splits = [
    ("TẬP TRAIN (60%)", "2.286 mẫu hạt", "Học cây, lập bảng ccp_alpha, 5-Fold Cross Validation.", C_GREEN_BG, C_UTEHY_GREEN),
    ("TẬP VALIDATION (20%)", "762 mẫu hạt", "Dò tìm siêu tham số tối ưu, so sánh đối chứng mô hình.", C_BLUE_LIGHT, C_NAVY_DARK),
    ("TẬP TEST (20%)", "762 mẫu hạt", "Niêm phong tuyệt đối, chỉ nạp 1 lần duy nhất nghiệm thu.", C_RED_BG, C_RED)
]
for title_s, num_s, desc_s, bg_c, text_c in flow_splits:
    p_t = tf_s5_1.add_paragraph()
    p_t.text = f"• {title_s}: {num_s}"
    p_t.font.bold = True
    p_t.font.size = Pt(11)
    p_t.font.color.rgb = text_c
    p_t.space_before = Pt(8)
    p_d = tf_s5_1.add_paragraph()
    p_d.text = f"  {desc_s}"
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = C_DARK
    p_d.space_before = Pt(2)

p_leak = tf_s5_1.add_paragraph()
p_leak.text = "✓ Nguyên tắc chống Leakage: Mọi ngưỡng chuẩn hóa và min-max chỉ tính từ tập Train!"
p_leak.font.bold = True
p_leak.font.size = Pt(10)
p_leak.font.color.rgb = C_UTEHY_GREEN
p_leak.space_before = Pt(12)

# Cột 2: Cột mốc Dummy Baseline
add_card(s5, Inches(4.84), Inches(1.15), Inches(3.64), Inches(5.55), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_s5_2 = s5.shapes.add_textbox(Inches(5.04), Inches(1.3), Inches(3.24), Inches(5.25))
tf_s5_2 = tb_s5_2.text_frame
tf_s5_2.word_wrap = True
p = tf_s5_2.paragraphs[0]
p.text = "2. CỘT MỐC DUMMY BASELINE"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN

p_desc = tf_s5_2.add_paragraph()
p_desc.text = "• Nguyên tắc: Không bao giờ đánh giá mô hình trong sự cô lập thiếu điểm tựa đối chứng."
p_desc.font.size = Pt(10)
p_desc.font.color.rgb = C_DARK
p_desc.space_before = Pt(6)

p_strat = tf_s5_2.add_paragraph()
p_strat.text = "• Chiến lược: DummyClassifier(strategy='most_frequent') luôn luôn dự đoán nhãn Osmancik."
p_strat.font.size = Pt(10)
p_strat.font.color.rgb = C_DARK
p_strat.space_before = Pt(4)

# KPI Box Dummy
add_kpi_badge(s5, Inches(5.04), Inches(3.2), Inches(3.24), Inches(1.6),
              "57,22%", "BASELINE ACCURACY\n(Lớp đa số Osmancik)",
              bg_color=C_RED_BG, num_color=C_RED, border_color=C_RED)

p_f1 = tf_s5_2.add_paragraph()
p_f1.text = "• F1-Score (Cammeo) = 0,0000 | ROC-AUC = 0.5000"
p_f1.font.bold = True
p_f1.font.size = Pt(10.5)
p_f1.font.color.rgb = C_NAVY_DARK
p_f1.space_before = Pt(80)

p_rule = tf_s5_2.add_paragraph()
p_rule.text = "✓ Thước đo thành công: Mô hình đề xuất bắt buộc phải vượt xa mốc 57,22%!"
p_rule.font.bold = True
p_rule.font.size = Pt(10)
p_rule.font.color.rgb = C_UTEHY_GREEN
p_rule.space_before = Pt(6)

# Cột 3: Trình tự 4 thí nghiệm thực chứng
add_card(s5, Inches(8.88), Inches(1.15), Inches(3.64), Inches(5.55), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_s5_3 = s5.shapes.add_textbox(Inches(9.08), Inches(1.3), Inches(3.24), Inches(5.25))
tf_s5_3 = tb_s5_3.text_frame
tf_s5_3.word_wrap = True
p = tf_s5_3.paragraphs[0]
p.text = "3. QUY TRÌNH 4 THÍ NGHIỆM"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN

experiments = [
    ("TN1: Khảo sát max_depth", "Quét depth 1 đến 20 chứng minh hiện tượng Overfitting của cây tự do."),
    ("TN2: Cắt tỉa ccp_alpha", "Rút gọn cấu trúc cây qua 5-Fold Cross Validation trên tập Train."),
    ("TN3: Cây đơn vs Rừng", "Đối chiếu toàn diện hiệu năng, ROC-AUC và độ minh bạch giải thích."),
    ("TN4: Nghiệm thu & Lỗi", "Kiểm thử Test Set 1 lần (91,60%) và mổ xẻ 50 ca lỗi trên tập Val.")
]
for title_e, desc_e in experiments:
    p_e = tf_s5_3.add_paragraph()
    p_e.text = f"• {title_e}"
    p_e.font.bold = True
    p_e.font.size = Pt(10.5)
    p_e.font.color.rgb = C_NAVY_DARK
    p_e.space_before = Pt(6)
    p_ed = tf_s5_3.add_paragraph()
    p_ed.text = f"  {desc_e}"
    p_ed.font.size = Pt(9.5)
    p_ed.font.color.rgb = C_DARK
    p_ed.space_before = Pt(2)

set_speaker_notes(s5, """
Để chống rò rỉ dữ liệu, nhóm phân chia phân tầng Stratified 60% Train, 20% Val và 20% Test, trong đó tập Test được niêm phong tuyệt đối. Chúng em cũng xây dựng Dummy Baseline luôn đoán nhãn đa số đạt sẵn 57,22% làm cột mốc sàn. Mọi thuật toán đề xuất bắt buộc phải chứng minh vượt xa cột mốc này qua 4 thí nghiệm thực chứng kế tiếp.
""")


# ============================================================
# SLIDE 6: THÍ NGHIỆM 1 - KHẢO SÁT OVERFITTING (KPI Cards + Hình 3)
# ============================================================
s6 = create_content_slide()
add_slide_header(s6, "Cây Quyết Định CART Bị Quá Khớp Nghiêm Trọng Khi Tăng Độ Sâu")
add_slide_footer(s6, 6)

# Cột trái: 3 Khối KPI / Stat Cards làm nổi bật hiện tượng Overfitting
add_kpi_badge(s6, Inches(0.8), Inches(1.15), Inches(5.7), Inches(1.65),
              "100,0% TRAIN ACCURACY",
              "BÙNG NỔ HỌC VẸT KHI MAX_DEPTH >= 16\nCây tự do ghi nhớ toàn bộ từng mẫu nhiễu và ngoại lai (High Variance)",
              bg_color=C_RED_BG, num_color=C_RED, border_color=C_RED)

add_kpi_badge(s6, Inches(0.8), Inches(2.95), Inches(5.7), Inches(1.85),
              "93,70% → 89,10% VAL ACC",
              "SUY THOÁI HIỆU NĂNG TRÊN TẬP VALIDATION\nĐạt đỉnh 93,70% ở depth=1, 2; giảm xuống đáy 89,10% ở depth=17 và ~89,90% ở depth=20",
              bg_color=C_GOLD_BG, num_color=C_NAVY_DARK, border_color=C_CARD_BORDER)

add_kpi_badge(s6, Inches(0.8), Inches(4.95), Inches(5.7), Inches(1.75),
              "GAP = 10,90% (PHƯƠNG SAI CAO)",
              "KẾT LUẬN THÍ NGHIỆM 1:\nCây quyết định tự do không có cơ chế tự dừng; bắt buộc phải cắt tỉa để tổng quát hóa!",
              bg_color=C_WHITE, num_color=C_UTEHY_GREEN, border_color=C_CARD_BORDER)

# Cột phải: Hình 3 Đường cong max_depth
img_p = os.path.join(FIGURES_DIR, "05_depth_curve.png")
if os.path.exists(img_p):
    s6.shapes.add_picture(img_p, Inches(6.75), Inches(1.15), width=Inches(5.78))
    tb_cap = s6.shapes.add_textbox(Inches(6.75), Inches(6.25), Inches(5.78), Inches(0.35))
    tb_cap.text_frame.paragraphs[0].text = "Hình 3: Đường cong hiệu năng: Val đạt đỉnh 93,70% (depth 1, 2) rồi giảm xuống 89,10% (depth 17) và ~89,90% (depth 20)"
    tb_cap.text_frame.paragraphs[0].font.size = Pt(9.5)
    tb_cap.text_frame.paragraphs[0].font.italic = True
    tb_cap.text_frame.paragraphs[0].font.color.rgb = C_MUTED
    tb_cap.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

set_speaker_notes(s6, """
Ở Thí nghiệm 1, khi tăng dần max_depth, đường màu xanh (Train) tăng vọt lên 100% khi depth đạt 16, nghĩa là cây đã học vẹt toàn bộ dữ liệu. Ngược lại, đường màu đỏ (Validation) chỉ đạt đỉnh 93,70% ở depth 1 và 2, sau đó giảm dần xuống đáy 89,10% ở depth 17 và dao động quanh 89,90% ở depth 20. Sự phân kỳ hơn 10% giữa Train và Val chứng minh cây tự do bị quá khớp nặng, khẳng định bắt buộc phải cắt tỉa mô hình.
""")


# ============================================================
# SLIDE 7: THÍ NGHIỆM 2 - CẮT TỈA CCP_ALPHA (Sơ đồ luật trực quan + Hình 4)
# ============================================================
s7 = create_content_slide()
add_slide_header(s7, "Cắt Tỉa ccp_alpha Giảm Hơn 97% Kích Thước Nhưng Giữ Vững Hiệu Năng")
add_slide_footer(s7, 7)

# Cột trái: 3 Khối Card tóm lược kết quả cắt tỉa + Sơ đồ cây luật
# Top KPI Card: 3 chỉ số then chốt
add_card(s7, Inches(0.8), Inches(1.15), Inches(5.7), Inches(2.2), bg_color=C_GREEN_BG, border_color=C_UTEHY_GREEN)
tb_s7_kpi = s7.shapes.add_textbox(Inches(1.0), Inches(1.25), Inches(5.3), Inches(2.0))
tf_s7_kpi = tb_s7_kpi.text_frame
tf_s7_kpi.word_wrap = True
p = tf_s7_kpi.paragraphs[0]
p.text = "KẾT QUẢ ĐỘT PHÁ CỦA CẮT TỈA CCP_ALPHA"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN

kpi_items_s7 = [
    ("Rút gọn hơn 97% kích thước:", "Từ 187 nút lá cồng kềnh xuống đúng 3 nút lá tinh gọn (tổng 5 nút)."),
    ("Bảo toàn hiệu năng Validation:", "Accuracy vẫn đạt 93,70% và F1-Score đạt 92,77%."),
    ("Tối ưu bằng 5-Fold CV:", "Xác định alpha* = 0.008024 với CV Accuracy = 92,48% ± 1,59%.")
]
for k_title, k_desc in kpi_items_s7:
    p_k = tf_s7_kpi.add_paragraph()
    p_k.text = f"• {k_title} "
    p_k.font.bold = True
    p_k.font.size = Pt(11)
    p_k.font.color.rgb = C_NAVY_DARK
    p_k.space_before = Pt(4)
    run = p_k.add_run()
    run.text = k_desc
    run.font.bold = False
    run.font.size = Pt(10.5)
    run.font.color.rgb = C_DARK

# Bottom Card: Sơ đồ luật quyết định 3 lá minh bạch (Chỉ dùng 2 biến)
add_card(s7, Inches(0.8), Inches(3.5), Inches(5.7), Inches(3.2), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_rule = s7.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(5.3), Inches(3.0))
tf_rule = tb_rule.text_frame
tf_rule.word_wrap = True

p = tf_rule.paragraphs[0]
p.text = "SƠ ĐỒ QUY TẮC QUYẾT ĐỊNH (CHỈ DÙNG 2 ĐẶC TRƯNG)"
p.font.size = Pt(11.5)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN

rules = [
    ("Luật 1 (Lá 1):", "Nếu Major_Axis_Length <= 191.05 px  -->  Dự đoán: OSMANCIK"),
    ("Luật 2 (Lá 2):", "Nếu Major_Axis_Length > 191.05 px VÀ Perimeter <= 467.86 px  -->  CAMMEO"),
    ("Luật 3 (Lá 3):", "Nếu Major_Axis_Length > 191.05 px VÀ Perimeter > 467.86 px  -->  CAMMEO")
]
for r_name, r_logic in rules:
    p_r = tf_rule.add_paragraph()
    p_r.text = f"{r_name} "
    p_r.font.bold = True
    p_r.font.size = Pt(10.5)
    p_r.font.color.rgb = C_NAVY_DARK
    p_r.space_before = Pt(5)
    run = p_r.add_run()
    run.text = r_logic
    run.font.bold = False
    run.font.size = Pt(10)
    run.font.color.rgb = C_DARK

p_note = tf_rule.add_paragraph()
p_note.text = "✓ 100% Minh bạch giải thích: Chỉ dùng Major_Axis_Length và Perimeter (Hoàn toàn không có Eccentricity). Người vận hành có thể nhớ và kiểm toán luật trong 5 giây!"
p_note.font.bold = True
p_note.font.size = Pt(9.5)
p_note.font.color.rgb = C_UTEHY_GREEN
p_note.space_before = Pt(8)

# Cột phải: Hình 4 Sơ đồ cây quyết định
img_p = os.path.join(FIGURES_DIR, "07_decision_tree.png")
if os.path.exists(img_p):
    s7.shapes.add_picture(img_p, Inches(6.75), Inches(1.15), width=Inches(5.78))
    tb_cap = s7.shapes.add_textbox(Inches(6.75), Inches(6.25), Inches(5.78), Inches(0.35))
    tb_cap.text_frame.paragraphs[0].text = "Hình 4: Sơ đồ Cây sau cắt tỉa: chỉ gồm 3 nút lá tinh gọn, sử dụng Major_Axis_Length và Perimeter"
    tb_cap.text_frame.paragraphs[0].font.size = Pt(9.5)
    tb_cap.text_frame.paragraphs[0].font.italic = True
    tb_cap.text_frame.paragraphs[0].font.color.rgb = C_MUTED
    tb_cap.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

set_speaker_notes(s7, """
Tại Thí nghiệm 2, nhóm áp dụng kỹ thuật Cost-Complexity Pruning qua 5-Fold Cross Validation trên tập Train để tìm ra alpha tối ưu bằng 0.008024. Kết quả cực kỳ ấn tượng: từ một cây cồng kềnh gồm 187 lá, mô hình được rút gọn chỉ còn đúng 3 nút lá (tổng cộng 5 nút). Quan sát sơ đồ cây ở Hình 4, cây chỉ cần đúng 2 biến là Major_Axis_Length và Perimeter để ra quyết định, hoàn toàn không cần đến Eccentricity. Độ chính xác trên Validation vẫn đạt 93,70%, mang lại tính minh bạch tuyệt đối cho người vận hành.
""")


# ============================================================
# SLIDE 8: THÍ NGHIỆM 3 - BẢNG SO SÁNH CHUẨN & ROC-AUC (Table + Hình 5)
# ============================================================
s8 = create_content_slide()
add_slide_header(s8, "Cây Cắt Tỉa Vượt Trội Về Minh Bạch, Rừng Ngẫu Nhiên Tối Ưu ROC-AUC")
add_slide_footer(s8, 8)

# Bảng so sánh mô hình chuẩn PPTX bên trái (Left: 0.8 in, Width: 5.75 in)
tb_w = Inches(5.75)
tb_h = Inches(3.2)
table_shape = s8.shapes.add_table(5, 6, Inches(0.8), Inches(1.15), tb_w, tb_h)
tbl = table_shape.table

col_widths = [Inches(1.85), Inches(0.75), Inches(0.75), Inches(0.95), Inches(0.75), Inches(0.70)]
for idx, width in enumerate(col_widths):
    tbl.columns[idx].width = width

table_data = [
    ["Mô hình", "Val Acc", "Val F1", "Val ROC-AUC", "Thời gian", "Minh bạch"],
    ["Dummy Baseline", "57,22%", "0,00%", "0.5000", "< 0.001s", "Đoán đa số"],
    ["Decision Tree (Gốc)", "89,24%", "87,41%", "0.8901", "0.015s", "Thấp (187 lá)"],
    ["Decision Tree (Đã tỉa)", "93,70%", "92,77%", "0.9596 (~0.96)", "0.036s", "Cực cao (3 lá)"],
    ["Random Forest (100 cây)", "93,44%", "92,28%", "0.9741 (~0.97)", "0.320s", "Hộp đen"]
]

for row_idx, row_values in enumerate(table_data):
    for col_idx, cell_value in enumerate(row_values):
        cell = tbl.cell(row_idx, col_idx)
        cell.text = cell_value
        p = cell.text_frame.paragraphs[0]
        p.font.name = "Calibri"
        
        if row_idx == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_UTEHY_GREEN
            p.font.bold = True
            p.font.size = Pt(10)
            p.font.color.rgb = C_WHITE
            p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT
        else:
            p.font.size = Pt(9.5)
            p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT
            if row_idx == 3: # DT Pruned Highlight
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_GREEN_BG
                p.font.bold = True
                p.font.color.rgb = C_NAVY_DARK
            elif row_idx == 4: # RF Highlight
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_BLUE_LIGHT
                p.font.bold = True
                p.font.color.rgb = C_NAVY_DARK
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_WHITE
                p.font.color.rgb = C_DARK

# Card phân tích Trade-off bên dưới Bảng
add_card(s8, Inches(0.8), Inches(4.5), Inches(5.75), Inches(2.2), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_insight = s8.shapes.add_textbox(Inches(1.0), Inches(4.6), Inches(5.35), Inches(2.0))
tf_in = tb_insight.text_frame
tf_in.word_wrap = True

p_in = tf_in.paragraphs[0]
p_in.text = "NHẬN XÉT ĐỐI CHỨNG VÀ ĐỒNG BỘ CHỈ SỐ:"
p_in.font.size = Pt(11)
p_in.font.bold = True
p_in.font.color.rgb = C_UTEHY_GREEN

insights_s8 = [
    ("Đồng bộ chỉ số ROC-AUC:", "DT Đã tỉa đạt Val ROC-AUC = 0.9596 (làm tròn 0.96 trên Hình 5), trong khi Random Forest đạt 0.9741 (làm tròn 0.97 trên Hình 5)."),
    ("Hiệu năng Accuracy tương đương:", "Cây cắt tỉa (93,70%) nhỉnh hơn nhẹ so với Rừng ngẫu nhiên (93,44%) trên tập Validation."),
    ("Đánh đổi cốt lõi (Trade-off):", "Cây tỉa tối ưu về tốc độ (0.036s) và minh bạch 100% để giải thích cho công nhân; Rừng ngẫu nhiên tối ưu đường cong xác suất liên tục (AUC 0.97) giúp phân loại van khí nén an toàn.")
]
for b_title, b_desc in insights_s8:
    p = tf_in.add_paragraph()
    p.text = f"• {b_title} "
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = C_NAVY_DARK
    p.space_before = Pt(3)
    run = p.add_run()
    run.text = b_desc
    run.font.bold = False
    run.font.size = Pt(9.5)
    run.font.color.rgb = C_DARK

# Hình 5 ROC Curves bên phải
img_p = os.path.join(FIGURES_DIR, "11_roc_curves.png")
if os.path.exists(img_p):
    s8.shapes.add_picture(img_p, Inches(6.75), Inches(1.15), width=Inches(5.78))
    tb_cap = s8.shapes.add_textbox(Inches(6.75), Inches(6.25), Inches(5.78), Inches(0.35))
    tb_cap.text_frame.paragraphs[0].text = "Hình 5: Đường cong ROC: Decision Tree Đã tỉa (AUC = 0.96) và Random Forest (AUC = 0.97)"
    tb_cap.text_frame.paragraphs[0].font.size = Pt(9.5)
    tb_cap.text_frame.paragraphs[0].font.italic = True
    tb_cap.text_frame.paragraphs[0].font.color.rgb = C_MUTED
    tb_cap.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

set_speaker_notes(s8, """
Bảng so sánh ở Thí nghiệm 3 tổng hợp đầy đủ các tiêu chí cốt lõi: Cả Cây quyết định đã tỉa và Rừng ngẫu nhiên đều vượt xa Baseline 57,22%. Về Accuracy, Cây đã tỉa đạt 93,70%, tương đương với Rừng 93,44%. Về chỉ số ROC-AUC, DT Đã tỉa đạt 0.9596 (làm tròn 0.96 trên Hình 5), còn Rừng ngẫu nhiên đạt 0.9741 (làm tròn 0.97). Bài toán thể hiện sự đánh đổi rõ ràng: Cây 3 lá cực kỳ phù hợp cho khâu kiểm toán và giải thích nghiệp vụ, còn Rừng ngẫu nhiên tối ưu về xác suất tin cậy để điều khiển van phân loại tự động.
""")


# ============================================================
# SLIDE 9: THÍ NGHIỆM 4 - NGHIỆM THU TEST & PHÂN TÍCH LỖI (KPI Cards + Hình 6)
# ============================================================
s9 = create_content_slide()
add_slide_header(s9, "Nghiệm Thu Tập Test Đạt 91,60% Và Bản Chất Sai Số Ranh Giới Tự Nhiên")
add_slide_footer(s9, 9)

# Cột trái: 3 Khối Card phân tích ngắn gọn, trực quan
add_kpi_badge(s9, Inches(0.8), Inches(1.15), Inches(5.7), Inches(1.65),
              "91,60% TEST ACCURACY",
              "NGHIỆM THU TẬP TEST (1 LẦN DUY NHẤT)\nTest F1 = 89,76% | Test ROC-AUC = 0.9709. Độ lệch Val-Test chỉ 1,84% khẳng định KHÔNG rò rỉ dữ liệu!",
              bg_color=C_GREEN_BG, num_color=C_UTEHY_GREEN, border_color=C_UTEHY_GREEN)

add_card(s9, Inches(0.8), Inches(2.95), Inches(5.7), Inches(1.85), bg_color=C_WHITE, border_color=C_CARD_BORDER)
tb_err = s9.shapes.add_textbox(Inches(1.0), Inches(3.05), Inches(5.3), Inches(1.65))
tf_err = tb_err.text_frame
tf_err.word_wrap = True
p = tf_err.paragraphs[0]
p.text = "PHÂN TÍCH 50 CA ĐOÁN SAI TRÊN TẬP VALIDATION"
p.font.size = Pt(11.5)
p.font.bold = True
p.font.color.rgb = C_RED

err_bullets = [
    ("Quy tắc khoa học:", "Trích xuất từ 762 mẫu Val của RF (762 x 6,56% = 50 mẫu), không dò xét tập Test."),
    ("Vùng giao thoa tự nhiên:", "Mẫu sai tập trung ở Major_Axis_Length từ 180 đến 205 px (quanh ngưỡng 191 px)."),
    ("Bản chất sinh học:", "Hạt Osmancik to bất thường xen lẫn hạt Cammeo nhỏ còi.")
]
for b_t, b_d in err_bullets:
    p_b = tf_err.add_paragraph()
    p_b.text = f"• {b_t} "
    p_b.font.bold = True
    p_b.font.size = Pt(10)
    p_b.font.color.rgb = C_NAVY_DARK
    p_b.space_before = Pt(3)
    run = p_b.add_run()
    run.text = b_d
    run.font.bold = False
    run.font.size = Pt(9.5)
    run.font.color.rgb = C_DARK

add_kpi_badge(s9, Inches(0.8), Inches(4.95), Inches(5.7), Inches(1.75),
              "94,37% → 78,28% TỰ TIN",
              "NĂNG LỰC TỰ NHẬN THỨC BẤT ĐỊNH\nĐộ tự tin tự động sụt giảm mạnh khi gặp mẫu khó, cho phép phát cờ cảnh báo an toàn!",
              bg_color=C_BLUE_LIGHT, num_color=C_NAVY_DARK, border_color=C_CARD_BORDER)

# Cột phải: Hình 6 Biểu đồ phân tích lỗi
img_p = os.path.join(FIGURES_DIR, "12_error_analysis.png")
if os.path.exists(img_p):
    s9.shapes.add_picture(img_p, Inches(6.75), Inches(1.15), width=Inches(5.78))
    tb_cap = s9.shapes.add_textbox(Inches(6.75), Inches(6.25), Inches(5.78), Inches(0.35))
    tb_cap.text_frame.paragraphs[0].text = "Hình 6: Phân tích 50 ca sai trên tập Val của RF: tập trung tại vùng giao thoa Major_Axis_Length (180 - 205 px)"
    tb_cap.text_frame.paragraphs[0].font.size = Pt(9.5)
    tb_cap.text_frame.paragraphs[0].font.italic = True
    tb_cap.text_frame.paragraphs[0].font.color.rgb = C_MUTED
    tb_cap.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

set_speaker_notes(s9, """
Ở bước nghiệm thu cuối cùng, mô hình Random Forest được đánh giá một lần duy nhất trên tập Test độc lập và đạt Accuracy 91,60%, ROC-AUC 0.9709. Độ chênh lệch giữa Val và Test chỉ 1,84%, chứng minh mô hình tổng quát hóa xuất sắc và quy trình không bị Data Leakage. Để hiểu rõ nguyên nhân sai số mà không vi phạm nguyên tắc niêm phong tập Test, nhóm mổ xẻ 50 ca đoán sai trên tập Validation. Kết quả ở Hình 6 khẳng định các mẫu sai màu đỏ tập trung tại vùng giao thoa tự nhiên Major_Axis_Length từ 180 đến 205 px (quanh ngưỡng phân chia 191 px). Điều đặc biệt là độ tự tin của mô hình tự động giảm từ 94% xuống 78% ở các ca này, mở ra cơ chế cảnh báo bất định.
""")


# ============================================================
# SLIDE 10: ỨNG DỤNG WEB APP THỰC TẾ (Sơ đồ kiến trúc 3 khối)
# ============================================================
s10 = create_content_slide()
add_slide_header(s10, "Đóng Gói Web App Flask Với Kiểm Soát 3 Tầng Và Nhận Diện Bất Định")
add_slide_footer(s10, 10)

cards_s10 = [
    (Inches(0.8), Inches(1.15), Inches(3.64), Inches(5.55), "1. TÁCH BIỆT SERVING OFFLINE", [
        ("Nạp mô hình tĩnh:", "final_model.pkl nạp vào RAM 1 lần duy nhất khi khởi động."),
        ("Độ trễ suy luận:", "< 1 ms / hạt, sẵn sàng tích hợp băng chuyền công nghiệp."),
        ("Giao diện chuẩn hóa:", "7 tham số tiếng Việt kèm đơn vị đo chuẩn (pixel, pixel²)."),
        ("Nạp nhanh mẫu thử:", "2 nút điền dữ liệu mẫu thực tế cho Cammeo và Osmancik.")
    ], C_WHITE, C_UTEHY_GREEN),
    (Inches(4.84), Inches(1.15), Inches(3.64), Inches(5.55), "2. BẢO VỆ ĐẦU VÀO 3 TẦNG", [
        ("Tầng 1 (Toàn vẹn):", "Chặn trường rỗng hoặc thiếu tham số (HTTP 400)."),
        ("Tầng 2 (Kiểu số):", "Bắt lỗi ký tự lạ, chữ cái hoặc giá trị âm vô lý."),
        ("Tầng 3 (Miền giá trị):", "So khớp khoảng [min, max] tập Train (Cảnh báo vàng)."),
        ("Trách nhiệm dữ liệu:", "Tuyệt đối không tự ý thay đổi dữ liệu của người dùng.")
    ], C_WHITE, C_RED),
    (Inches(8.88), Inches(1.15), Inches(3.64), Inches(5.55), "3. TÍNH NĂNG ĐẠT ĐIỂM CỘNG", [
        ("Cờ bất định (Uncertainty):", "Dải xác suất [0.40, 0.60] tự động báo cờ đỏ chuyển soi kính."),
        ("Dự đoán theo lô (Batch):", "Upload file CSV hàng trăm hạt, xuất kết quả tức thì."),
        ("Công bố Model Card:", "Minh bạch phiên bản scikit-learn 1.4.2 & đạo đức AI."),
        ("Bản quyền dữ liệu:", "Công nhận chuẩn UCI Machine Learning Repository.")
    ], C_WHITE, C_UTEHY_GOLD)
]

for left, top, w, h, title, items, bg_c, border_c in cards_s10:
    add_card(s10, left, top, w, h, bg_color=bg_c, border_color=border_c)
    tb = s10.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), w - Inches(0.4), h - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = border_c
    for b_title, b_desc in items:
        p_item = tf.add_paragraph()
        p_item.text = f"• {b_title} "
        p_item.font.bold = True
        p_item.font.size = Pt(10.5)
        p_item.font.color.rgb = C_NAVY_DARK
        p_item.space_before = Pt(6)
        run = p_item.add_run()
        run.text = b_desc
        run.font.bold = False
        run.font.size = Pt(10)
        run.font.color.rgb = C_DARK

set_speaker_notes(s10, """
Nhóm đã đóng gói toàn bộ mô hình thành ứng dụng Web Flask hoàn chỉnh. Điểm nổi bật là kiến trúc tách biệt Offline Serving tải mô hình tĩnh 1 lần vào RAM giúp suy luận tức thì dưới 1 mili-giây. Hệ thống áp dụng cơ chế kiểm soát đầu vào 3 tầng nghiêm ngặt và đặc biệt là tính năng Cảnh báo bất định: khi xác suất rơi vào khoảng 0.40 đến 0.60, hệ thống tự động phát cảnh báo yêu cầu chuyên viên soi kính hiển vi thủ công, bảo vệ an toàn thương hiệu cho các lô hàng xuất khẩu.
""")


# ============================================================
# SLIDE 11: TỔNG KẾT ĐÓNG GÓP & HƯỚNG PHÁT TRIỂN (Visual Grid Cards)
# ============================================================
s11 = create_content_slide()
add_slide_header(s11, "Tổng Kết Đóng Góp Đề Tài Và Định Hướng Phát Triển Tương Lai")
add_slide_footer(s11, 11)

# Cột trái: 3 Đóng góp cốt lõi (Green Theme)
add_card(s11, Inches(0.8), Inches(1.15), Inches(5.7), Inches(5.55), bg_color=C_GREEN_BG, border_color=C_UTEHY_GREEN)
tb_s11_l = s11.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(5.3), Inches(5.25))
tf_s11_l = tb_s11_l.text_frame
tf_s11_l.word_wrap = True
p = tf_s11_l.paragraphs[0]
p.text = "3 ĐÓNG GÓP CỐT LÕI CỦA ĐỀ TÀI"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_NAVY_DARK

contribs = [
    ("(1) Tinh giản cây bằng ccp_alpha:", "Giảm hơn 97% kích thước (từ 187 lá xuống 3 nút lá), chỉ dựa vào 2 biến Major_Axis_Length và Perimeter mà vẫn đạt Validation Accuracy 93,70%."),
    ("(2) Random Forest & Nghiệm thu Test:", "Đạt ROC-AUC 0.97 và Test Accuracy 91,60% (vượt xa Baseline 57,22%), khẳng định quy trình Stratified chuẩn hóa không bị rò rỉ dữ liệu."),
    ("(3) Đóng gói Web Service an toàn:", "Triển khai hoàn chỉnh ứng dụng Flask với kiểm soát dữ liệu 3 tầng và cờ cảnh báo bất định [0.40, 0.60] bảo vệ chất lượng lô gạo.")
]
for c_title, c_desc in contribs:
    p_c = tf_s11_l.add_paragraph()
    p_c.text = f"{c_title}"
    p_c.font.bold = True
    p_c.font.size = Pt(11)
    p_c.font.color.rgb = C_UTEHY_GREEN
    p_c.space_before = Pt(8)
    p_cd = tf_s11_l.add_paragraph()
    p_cd.text = f"{c_desc}"
    p_cd.font.size = Pt(10)
    p_cd.font.color.rgb = C_DARK
    p_cd.space_before = Pt(2)

# Cột phải: 4 Hướng phát triển tương lai (Gold Theme)
add_card(s11, Inches(6.833), Inches(1.15), Inches(5.7), Inches(5.55), bg_color=C_WHITE, border_color=C_UTEHY_GOLD)
tb_s11_r = s11.shapes.add_textbox(Inches(7.033), Inches(1.3), Inches(5.3), Inches(5.25))
tf_s11_r = tb_s11_r.text_frame
tf_s11_r.word_wrap = True
p = tf_s11_r.paragraphs[0]
p.text = "ĐỊNH HƯỚNG PHÁT TRIỂN & MỞ RỘNG TƯƠNG LAI"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_NAVY_DARK

futures = [
    ("• Màu sắc & Kết cấu (RGB & GLCM):", "Khai thác thêm kênh màu và vân hạt để phân tách triệt để 8% mẫu nằm ở vùng ranh giới 180–205 px."),
    ("• Camera công nghiệp thời gian thực:", "Kết nối trực tiếp camera tốc độ cao ghi hình hạt rơi trên phễu rung, điều khiển van xịt khí nén."),
    ("• Mở rộng giống lúa Việt Nam:", "Nâng cấp phân loại đa lớp (Multi-class) cho các giống gạo đặc sản: ST24, ST25, Jasmine 85."),
    ("• Ứng dụng Deep Learning (CNNs):", "Thử nghiệm MobileNetV3 phân loại trực tiếp từ ảnh chụp khay hạt gạo không cần trích xuất đặc trưng tay.")
]
for f_title, f_desc in futures:
    p_f = tf_s11_r.add_paragraph()
    p_f.text = f_title
    p_f.font.bold = True
    p_f.font.size = Pt(11)
    p_f.font.color.rgb = C_NAVY_DARK
    p_f.space_before = Pt(6)
    p_fd = tf_s11_r.add_paragraph()
    p_fd.text = f_desc
    p_fd.font.size = Pt(10)
    p_fd.font.color.rgb = C_DARK
    p_fd.space_before = Pt(2)

set_speaker_notes(s11, """
Tổng kết lại, đề tài đã đạt được 3 kết quả cốt lõi: Thứ nhất, cắt tỉa ccp_alpha giúp rút gọn 97% kích thước cây xuống còn 3 nút lá mà vẫn đạt 93,70% Accuracy; Thứ hai, Random Forest đạt ROC-AUC 0.97 và Test Accuracy 91,60%, vượt xa mốc Baseline 57,22%; Thứ ba, đóng gói thành công Web Service Flask với kiểm soát 3 tầng và cờ cảnh báo bất định. Về hướng phát triển, nhóm định hướng bổ sung đặc trưng màu sắc và kết cấu hạt để giải quyết triệt để 8% mẫu ở vùng ranh giới 180-205 px, kết nối camera công nghiệp thời gian thực và mở rộng sang các giống gạo Việt Nam như ST25.
""")


# ============================================================
# SLIDE 12: LỜI CẢM ƠN & PHIÊN HỎI ĐÁP (Q&A) (Both Students)
# ============================================================
s12 = create_content_slide()
add_slide_header(s12, "Ba Bài Học Thực Tiễn Và Phiên Thảo Luận Báo Cáo Bảo Vệ")
add_slide_footer(s12, 12)

# 3 Khối Card tóm tắt bài học cốt lõi ở nửa trên
coords_s12 = [
    (Inches(0.8), Inches(1.15), Inches(3.64), Inches(2.6), "1. QUY TRÌNH QUYẾT ĐỊNH ĐỘ TIN CẬY",
     "Chia tầng Stratified 60/20/20 và bảo vệ tập Test độc lập có giá trị khoa học vượt trội so với điểm số ảo bị rò rỉ dữ liệu (Data Leakage)."),
    (Inches(4.84), Inches(1.15), Inches(3.64), Inches(2.6), "2. ĐƠN GIẢN LÀ ĐỈNH CAO (OCCAM'S RAZOR)",
     "Cây 3 lá sau cắt tỉa ccp_alpha giải thích minh bạch 100% cho người vận hành mà hiệu năng vẫn tương đương Rừng ngẫu nhiên 100 cây."),
    (Inches(8.88), Inches(1.15), Inches(3.64), Inches(2.6), "3. TRÁCH NHIỆM TRONG CÔNG NGHỆ",
     "Một hệ thống AI ứng dụng phải biết nghi ngờ: tính năng cảnh báo bất định dải [0.40, 0.60] bảo vệ an toàn thương hiệu nông sản.")
]

for left, top, w, h, title, text in coords_s12:
    add_card(s12, left, top, w, h, bg_color=C_WHITE, border_color=C_UTEHY_GREEN)
    tb = s12.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), w - Inches(0.4), h - Inches(0.3))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_UTEHY_GREEN
    p_desc = tf.add_paragraph()
    p_desc.text = text
    p_desc.font.size = Pt(11)
    p_desc.font.color.rgb = C_DARK
    p_desc.space_before = Pt(6)

# Card Lời Cảm Ơn & Q&A Trang Trọng ở nửa dưới (Giữ trọn vẹn cả 2 sinh viên)
add_card(s12, Inches(0.8), Inches(3.95), Inches(11.733), Inches(2.75), bg_color=C_WHITE, border_color=C_UTEHY_GREEN)
tb_ty = s12.shapes.add_textbox(Inches(1.0), Inches(4.05), Inches(11.333), Inches(2.55))
tf_ty = tb_ty.text_frame
tf_ty.word_wrap = True

p = tf_ty.paragraphs[0]
p.text = "XIN TRÂN TRỌNG CẢM ƠN THẦY GIÁO VÀ HỘI ĐỒNG BẢO VỆ!"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = C_NAVY_DARK
p.alignment = PP_ALIGN.CENTER

p = tf_ty.add_paragraph()
p.text = "Em xin chân thành cảm ơn PGS.TS. Nguyễn Văn Hậu đã tận tình hướng dẫn và định hướng khoa học cho đề tài!"
p.font.size = Pt(13)
p.font.italic = True
p.font.color.rgb = C_DARK
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(4)

p = tf_ty.add_paragraph()
p.text = "PHIÊN BẢO VỆ & THẢO LUẬN (Q & A)"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GOLD
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(6)

p = tf_ty.add_paragraph()
p.text = "Sinh viên thực hiện: Chử Quốc Hào (MSV: 10123112)  —  Vũ Đình Hiếu (MSV: 10123131) | Lớp: 10123O.1"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = C_UTEHY_GREEN
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(4)

p = tf_ty.add_paragraph()
p.text = "Mã nguồn đề tài: https://github.com/hao0979/Rice_Classification | Báo cáo: BaoCao_BTL_PhanLoaiGiongLua_final.docx"
p.font.size = Pt(10)
p.font.color.rgb = C_MUTED
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(3)

set_speaker_notes(s12, """
Kính thưa Thầy giáo PGS.TS. Nguyễn Văn Hậu và quý Thầy Cô trong Hội đồng! Bài tập lớn kết thúc học phần này không chỉ giúp chúng em nắm vững các kỹ thuật cốt lõi của Học máy cơ bản mà còn mang lại 3 bài học thực tiễn vô cùng giá trị: Quy trình chống rò rỉ dữ liệu quyết định độ tin cậy; Cây tối giản 3 lá minh chứng cho nguyên lý đơn giản là đỉnh cao; và một hệ thống AI có trách nhiệm phải biết cảnh báo vùng bất định. Nhóm chúng em xin trân trọng cảm ơn Thầy Hậu đã tận tình hướng dẫn. Chúng em xin sẵn sàng lắng nghe các câu hỏi phản biện và góp ý của Thầy để hoàn thiện hơn nữa đề tài ạ!
""")


# ============================================================
# LƯU FILE PRESENTATION VÀO CẢ 2 ĐƯỜNG DẪN
# ============================================================
saved_count = 0
try:
    prs.save(SLIDES_PATH)
    print(f"[OK] Da luu thanh cong slide chinh: {SLIDES_PATH}")
    saved_count += 1
except Exception as e:
    print(f"[ERROR] Loi luu {SLIDES_PATH}: {e}")

try:
    prs.save(ALT_PATH)
    print(f"[OK] Da luu thanh cong slide phu: {ALT_PATH}")
    saved_count += 1
except Exception as e:
    print(f"[ERROR] Loi luu {ALT_PATH}: {e}")

if saved_count > 0:
    print("HOAN THANH: Da tao bo slide 12 trang truc quan, toi uu it chu va day du Speaker Notes!")
else:
    raise RuntimeError("Khong the luu slide vao bat ky duong dan nao!")
